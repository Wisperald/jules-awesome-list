"""Сборка Excel-версии инфографики «Ислам: большая хронология 570–2026».

Лист «Инфографика» — диаграмма-лента: столбцы-десятилетия, полосы рисуются условным
форматированием по формулам из столбцов «Начало»/«Конец». Лист «Сводка» — статистика
по разделам (формулы). Лист «О документе» — легенда и оговорки.

Запуск: python3 build_xlsx.py <data.json> <out.xlsx>
data.json получается из ../infographic/data.js (см. README).
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

data = json.load(open(sys.argv[1], encoding="utf-8"))
OUT = sys.argv[2]

FONT = "Arial"
COL = {  # цвета полос (проверенная на различимость палитра инфографики)
    "sun": "0E7A5A", "shi": "3346A8", "kha": "B4531F", "kal": "A06A00", "suf": "0090B0",
    "mod": "C0395A", "ext": "8B1E1E", "het": "6F52A8", "mix": "6B7280", "gold": "B8892B",
    "pol": "C0395A", "ev": "B8892B", "ev2": "8A5F10",
}
EM_D, EM, GOLD_L, PAPER, LINE = "062E23", "0E5A43", "E9C877", "FBF9F4", "E6DFCF"

Y0, Y1, STEP = 570, 2020, 10            # столбцы-десятилетия: 570, 580, …, 2020
YEARS = list(range(Y0, Y1 + 1, STEP))
FIRST_GRID_COL = 10                      # J
C_SEC, C_NAME, C_START, C_END, C_STATUS, C_DUR, C_ORIGIN, C_COLOR, C_LIVE = 1, 2, 3, 4, 5, 6, 7, 8, 9
LAST_GRID_COL = FIRST_GRID_COL + len(YEARS) - 1
END_COL = LAST_GRID_COL + 1              # столбец-маркер «▶»


def gcol(year):
    return FIRST_GRID_COL + (max(Y0, min(year, Y1 + STEP - 1)) - Y0) // STEP


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


wb = Workbook()
ws = wb.active
ws.title = "Инфографика"
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = EM

thin = Side(style="thin", color=LINE)

# ---------- ширины ----------
widths = {C_SEC: 20, C_NAME: 46, C_START: 8, C_END: 8, C_STATUS: 12, C_DUR: 7, C_ORIGIN: 30, C_COLOR: 6, C_LIVE: 3}
for c, w in widths.items():
    ws.column_dimensions[get_column_letter(c)].width = w
ws.column_dimensions[get_column_letter(C_COLOR)].hidden = True
for c in range(FIRST_GRID_COL, END_COL + 1):
    ws.column_dimensions[get_column_letter(c)].width = 1.7

# ---------- шапка ----------
def band(row, height, color):
    ws.row_dimensions[row].height = height
    for c in range(1, END_COL + 1):
        ws.cell(row, c).fill = fill(color)

band(1, 44, EM_D); band(2, 22, EM_D); band(3, 20, EM_D); band(4, 8, PAPER)
ws["A1"] = "ИСЛАМ · БОЛЬШАЯ ХРОНОЛОГИЯ 570 — 2026"
ws["A1"].font = Font(name=FONT, size=24, bold=True, color="FFFFFF")
ws["A2"] = ("Эпохи, события, державы, все ветви и расколы, школы права и вероучения, суфийские ордены, "
            "реформаторские и новые движения, личности. Столбец = десятилетие; полоса строится автоматически "
            "из столбцов «Начало» и «Конец».")
ws["A2"].font = Font(name=FONT, size=11, italic=True, color="F1E4C2")
ws["A3"] = ("Цвета:  ■ сунниты  ■ шииты  ■ хариджиты/ибадиты  ■ калам и философия  ■ суфизм  "
            "■ реформизм и новые движения  ■ гетеродоксия / спорный статус  ■ экстремизм (справочно)  "
            "·  ▶ — течение существует сегодня")
ws["A3"].font = Font(name=FONT, size=10, color="CFE3DA")
ws.cell(1, FIRST_GRID_COL + 118, "ٱلْإِسْلَام").font = Font(name=FONT, size=26, color=GOLD_L)

# ---------- эпохи ----------
R_ERA, R_AX, R_YR, R_HDR = 5, 6, 7, 8
ws.row_dimensions[R_ERA].height = 30
ws.cell(R_ERA, C_NAME, "ЭПОХИ").font = Font(name=FONT, size=10, bold=True, color=EM)
ws.cell(R_ERA, C_NAME).alignment = Alignment(horizontal="right", vertical="center")
prev_c1 = FIRST_GRID_COL - 1
for i, (a, b, name) in enumerate(data["ERAS"]):
    c0, c1 = max(gcol(a), prev_c1 + 1), max(gcol(b - 1), prev_c1 + 1)
    if b >= 2026:
        c1 = LAST_GRID_COL
    prev_c1 = c1
    if c1 > c0:
        ws.merge_cells(start_row=R_ERA, start_column=c0, end_row=R_ERA, end_column=c1)
    cell = ws.cell(R_ERA, c0, name)
    cell.fill = fill(EM if i % 2 else "0B4A38")
    cell.font = Font(name=FONT, size=8, bold=True, color="FFFFFF")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True, shrink_to_fit=False)
    for c in range(c0, c1 + 1):
        ws.cell(R_ERA, c).fill = fill(EM if i % 2 else "0B4A38")

# ---------- ось: подписи каждые 50 лет + служебная строка с годом столбца ----------
ws.row_dimensions[R_AX].height = 16
ws.row_dimensions[R_YR].height = 9
for j, y in enumerate(YEARS):
    c = FIRST_GRID_COL + j
    ws.cell(R_YR, c, y).font = Font(name=FONT, size=5, color="B8B0A0")
    ws.cell(R_YR, c).alignment = Alignment(horizontal="center", text_rotation=90)
    if y % 50 == 0 or y == Y0:
        span = 5 if y % 50 == 0 else 3
        c1 = min(c + span - 1, LAST_GRID_COL)
        ws.merge_cells(start_row=R_AX, start_column=c, end_row=R_AX, end_column=c1)
        cell = ws.cell(R_AX, c, str(y))
        cell.font = Font(name=FONT, size=9 if y % 100 == 0 else 7, bold=y % 100 == 0, color="14171B" if y % 100 == 0 else "6A7078")
        cell.alignment = Alignment(horizontal="left")
ws.cell(R_AX, END_COL, "▶").font = Font(name=FONT, size=8, color=COL["gold"])

heads = ["Раздел", "Название (течение, держава, событие, личность)", "Начало", "Конец", "Статус", "Лет", "Откололось от / примечание", "Цвет", ""]
for c, h in enumerate(heads, 1):
    cell = ws.cell(R_HDR, c, h)
    cell.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
    cell.fill = fill(EM_D)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
for c in range(FIRST_GRID_COL, END_COL + 1):
    ws.cell(R_HDR, c).fill = fill(EM_D)
ws.row_dimensions[R_HDR].height = 28

# ---------- строки ----------
rows = []  # (section, name, start, end, color, origin, live, is_event)
sections = []  # (title, description, color, first_row, last_row)

def add_section(title, desc, color, items):
    sections.append([title, desc, color, items])

ev = [(f"{t}", y, y, "ev2" if imp > 1 else "ev", "", y >= 2026, True) for y, t, imp in data["EVENTS"]]
add_section("Ключевые события", "Даты, изменившие ход истории уммы", "3A4048", ev)

st = [(n, a, b, k, "", bool(lv[0]) if lv else False, False) for a, b, n, k, *lv in data["STATES"]]
add_section("Державы и династии", "Цвет — вероисповедание правящей династии", EM, st)

for L in data["LANES"]:
    labels = {it[0]: it[3] for it in L["items"]}
    base = {"kha": "kha", "shi": "shi", "sun": "sun", "kal": "kal", "suf": "suf", "mod": "mod", "ext": "ext"}[L["id"]]
    items = []
    for id_, a, b, label, parent, flags in L["items"]:
        color = "ext" if "W" in flags else ("het" if "H" in flags else base)
        note = []
        if parent:
            note.append("← " + labels.get(parent, ""))
        if "P" in flags:
            note.append("политический ислам")
        if "H" in flags:
            note.append("спорный статус / вне ислама")
        if "W" in flags:
            note.append("запрещена (РК/РФ/ООН)")
        items.append((label, a, b, color, "; ".join(note), "L" in flags, False))
    add_section(L["name"], L["note"], COL[base], items)

pp = [(n, a, b, k, "", bool(lv[0]) if lv else False, False) for a, b, n, k, *lv in data["PEOPLE"]]
add_section("Личности", "Годы жизни. Цвет — направление или сфера деятельности", "3A4048", pp)

r = R_HDR + 1
first_data = r
sec_ranges = []
for title, desc, color, items in sections:
    # строка-заголовок раздела
    ws.row_dimensions[r].height = 22
    for c in range(1, END_COL + 1):
        ws.cell(r, c).fill = fill(color)
    ws.cell(r, C_SEC, "■ " + title.upper()).font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    ws.cell(r, C_NAME, desc).font = Font(name=FONT, size=8, italic=True, color="FFFFFF")
    ws.cell(r, C_NAME).alignment = Alignment(vertical="center")
    ws.cell(r, C_SEC).alignment = Alignment(vertical="center")
    r += 1
    s0 = r
    items = sorted(items, key=lambda t: (t[1], t[2]))
    for name, a, b, color, origin, live, is_event in items:
        end = 2026 if live else b
        ws.cell(r, C_SEC, title).font = Font(name=FONT, size=7, color="9AA0A6")
        ws.cell(r, C_NAME, name).font = Font(name=FONT, size=9, bold=is_event and color == "ev2")
        ws.cell(r, C_START, a).font = Font(name=FONT, size=9, color="0000FF")
        ws.cell(r, C_END, end).font = Font(name=FONT, size=9, color="0000FF")
        ws.cell(r, C_STATUS, f'=IF({get_column_letter(C_START)}{r}={get_column_letter(C_END)}{r},"событие",IF({get_column_letter(C_END)}{r}>=2026,IF({get_column_letter(C_SEC)}{r}="Личности","жив","существует"),"в прошлом"))').font = Font(name=FONT, size=8)
        ws.cell(r, C_DUR, f'=IF({get_column_letter(C_START)}{r}={get_column_letter(C_END)}{r},"",{get_column_letter(C_END)}{r}-{get_column_letter(C_START)}{r})').font = Font(name=FONT, size=8, color="6A7078")
        ws.cell(r, C_ORIGIN, origin).font = Font(name=FONT, size=8, color="6A7078")
        ws.cell(r, C_COLOR, color).font = Font(name=FONT, size=7, color="9AA0A6")
        live_formula = f'=IF(AND({get_column_letter(C_END)}{r}>=2026,{get_column_letter(C_START)}{r}<{get_column_letter(C_END)}{r}),"▶","")'
        ws.cell(r, END_COL, live_formula).font = Font(name=FONT, size=8, color=COL.get(color, "000000"))
        for c in range(1, C_LIVE + 1):
            ws.cell(r, c).border = Border(bottom=thin)
        ws.row_dimensions[r].height = 13
        r += 1
    sec_ranges.append((title, s0, r - 1))
last_data = r - 1

# ---------- условное форматирование: полосы ----------
gl, gr = get_column_letter(FIRST_GRID_COL), get_column_letter(LAST_GRID_COL)
S, E, K = get_column_letter(C_START), get_column_letter(C_END), get_column_letter(C_COLOR)
grid = f"{gl}{first_data}:{gr}{last_data}"
top = f"{gl}${R_YR}"
for key, hex_ in COL.items():
    formula = f'AND({top}<=${E}{first_data},{top}+{STEP}>${S}{first_data},${K}{first_data}="{key}")'
    ws.conditional_formatting.add(grid, FormulaRule(formula=[formula], fill=fill(hex_), stopIfTrue=True))
# фон сетки: вертикальные линии веков
for j, y in enumerate(YEARS):
    if y % 100 == 0:
        c = FIRST_GRID_COL + j
        for rr in range(first_data, last_data + 1):
            ws.cell(rr, c).border = Border(left=Side(style="thin", color="D9CFB8"))

ws.freeze_panes = ws.cell(R_HDR + 1, FIRST_GRID_COL)
ws.auto_filter.ref = f"A{R_HDR}:{get_column_letter(C_ORIGIN)}{last_data}"
ws.print_title_rows = f"{R_ERA}:{R_HDR}"
ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = ws.PAPERSIZE_A3
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_options.horizontalCentered = True

note_row = last_data + 2
ws.cell(note_row, C_SEC, "Примечания").font = Font(name=FONT, size=9, bold=True)
notes = [
    "Синие числа в «Начало»/«Конец» — вводимые данные: измените их, и полоса, статус и длительность пересчитаются автоматически.",
    "Столбец временной шкалы = десятилетие (570–579, 580–589 … 2020–2029); события внутри одного десятилетия попадают в одну ячейку. Точные годы — в столбцах C–D.",
    "Для существующих течений «Конец» = 2026 (состояние на осень 2026 г.). Даты ранних событий приблизительны.",
    "Источник: энциклопедия «Ислам: история веры, личности, течения» (2026) и её инфографика; численность — Pew Research Center (2025).",
    "Организации, отмеченные тёмно-красным, признаны террористическими в РК, РФ и/или ООН; сведения носят справочный характер.",
]
for i, t in enumerate(notes):
    ws.cell(note_row + 1 + i, C_SEC, t).font = Font(name=FONT, size=8, color="3A4048")

# ---------- лист «Сводка» ----------
sm = wb.create_sheet("Сводка")
sm.sheet_view.showGridLines = False
sm.sheet_properties.tabColor = COL["gold"]
sm["A1"] = "Сводка по разделам хронологии"
sm["A1"].font = Font(name=FONT, size=16, bold=True, color=EM_D)
sm["A2"] = "Все значения — формулы, ссылающиеся на лист «Инфографика»."
sm["A2"].font = Font(name=FONT, size=9, italic=True, color="6A7078")
hdr = ["Раздел", "Всего строк", "Существуют / живы сегодня", "В прошлом / события", "Самое раннее начало", "Самое позднее начало", "Доля существующих"]
for c, h in enumerate(hdr, 1):
    cell = sm.cell(4, c, h)
    cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    cell.fill = fill(EM_D)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
sm.row_dimensions[4].height = 30
src = "Инфографика"
rng = lambda col: f"{src}!${col}${first_data}:${col}${last_data}"
A_, C_, St_ = get_column_letter(C_SEC), get_column_letter(C_START), get_column_letter(C_STATUS)
for i, (title, s0, s1) in enumerate(sec_ranges):
    rr = 5 + i
    sm.cell(rr, 1, title).font = Font(name=FONT, size=10)
    sm.cell(rr, 2, f'=COUNTIF({rng(A_)},A{rr})')
    sm.cell(rr, 3, f'=COUNTIFS({rng(A_)},A{rr},{rng(St_)},"существует")+COUNTIFS({rng(A_)},A{rr},{rng(St_)},"жив")')
    sm.cell(rr, 4, f'=B{rr}-C{rr}')
    sm.cell(rr, 5, f'=_xlfn.MINIFS({rng(C_)},{rng(A_)},A{rr})')
    sm.cell(rr, 6, f'=_xlfn.MAXIFS({rng(C_)},{rng(A_)},A{rr})')
    sm.cell(rr, 7, f'=IF(B{rr}=0,"-",C{rr}/B{rr})')
    sm.cell(rr, 7).number_format = "0%"
    for c in range(2, 8):
        sm.cell(rr, c).font = Font(name=FONT, size=10)
tot = 5 + len(sec_ranges)
sm.cell(tot, 1, "ИТОГО").font = Font(name=FONT, size=10, bold=True)
for c, f in [(2, f"=SUM(B5:B{tot-1})"), (3, f"=SUM(C5:C{tot-1})"), (4, f"=SUM(D5:D{tot-1})"),
             (5, f"=MIN(E5:E{tot-1})"), (6, f"=MAX(F5:F{tot-1})"), (7, f'=IF(B{tot}=0,"-",C{tot}/B{tot})')]:
    sm.cell(tot, c, f).font = Font(name=FONT, size=10, bold=True)
    sm.cell(tot, c).border = Border(top=Side(style="medium", color=EM_D))
sm.cell(tot, 7).number_format = "0%"
sm.column_dimensions["A"].width = 36
for c in "BCDEFG":
    sm.column_dimensions[c].width = 16

k = tot + 3
sm.cell(k, 1, "Мусульманский мир сегодня (оценки)").font = Font(name=FONT, size=12, bold=True, color=EM_D)
facts = [("Мусульман в мире, млрд (2020)", 2.0), ("Доля в населении Земли", 0.256), ("Сунниты, доля (оценка 85–90%)", 0.875),
         ("Шииты, доля (оценка 10–13%)", 0.115), ("Ибадиты, доля", 0.002), ("Государств — членов ОИС", 57), ("Прогноз численности на 2050 г., млрд", 2.8)]
for i, (lab, v) in enumerate(facts):
    sm.cell(k + 1 + i, 1, lab).font = Font(name=FONT, size=10)
    cell = sm.cell(k + 1 + i, 2, v)
    cell.font = Font(name=FONT, size=10, color="0000FF")
    if isinstance(v, float) and v < 1:
        cell.number_format = "0.0%"
sm.cell(k + 1 + len(facts), 1, "Источник: Pew Research Center, «How the Global Religious Landscape Changed From 2010 to 2020» (2025); доли сунниты/шииты — середина диапазона оценок.").font = Font(name=FONT, size=8, italic=True, color="6A7078")

# ---------- лист «О документе» ----------
ab = wb.create_sheet("О документе")
ab.sheet_view.showGridLines = False
ab.column_dimensions["A"].width = 120
lines = [
    ("Ислам · Большая хронология 570–2026 — Excel-версия инфографики", Font(name=FONT, size=16, bold=True, color=EM_D)),
    ("", None),
    ("Как читать", Font(name=FONT, size=12, bold=True, color=EM)),
    ("• Лист «Инфографика»: каждая строка — событие, держава, течение или личность; цветная полоса справа показывает период существования.", None),
    ("• Столбец «Откололось от» показывает, от какой ветви произошло течение (точка раскола).", None),
    ("• ▶ в последнем столбце — течение или государство существует сегодня.", None),
    ("• Фильтры в заголовке таблицы позволяют отобрать раздел, статус или период.", None),
    ("", None),
    ("Как редактировать", Font(name=FONT, size=12, bold=True, color=EM)),
    ("• Меняйте только синие числа в столбцах «Начало» и «Конец» — полосы, статус, длительность и «Сводка» пересчитаются.", None),
    ("• Новую строку добавляйте внутри раздела: скопируйте соседнюю строку целиком (цветовой код в скрытом столбце H задаёт цвет полосы).", None),
    ("• Коды цветов: sun — сунниты, shi — шииты, kha — хариджиты/ибадиты, kal — калам и философия, suf — суфизм, mod — реформизм, het — гетеродоксия, ext — экстремизм, mix — смешанная, gold — умма, ev/ev2 — события.", None),
    ("", None),
    ("Оговорки", Font(name=FONT, size=12, bold=True, color=EM)),
    ("• Шкала — по десятилетиям; события одного десятилетия попадают в одну ячейку.", None),
    ("• Даты ранних периодов приблизительны; численность общин — диапазоны оценок.", None),
    ("• Сведения о 2025–2026 гг. — по состоянию на осень 2026 г.", None),
    ("• Упоминание запрещённых организаций носит справочно-исторический характер.", None),
]
for i, (t, f) in enumerate(lines, 1):
    ab.cell(i, 1, t).font = f or Font(name=FONT, size=10)
    ab.cell(i, 1).alignment = Alignment(wrap_text=True)

wb.save(OUT)
print("rows:", last_data - first_data + 1, "grid cols:", len(YEARS), "->", OUT)
